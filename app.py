
import streamlit as st
import pymupdf
import streamlit.components.v1 as components
from docx import Document
import json
import subprocess
import tempfile
import os
def read_doc_file(uploaded_file):
    with tempfile.TemporaryDirectory() as temp_dir:
        input_path = os.path.join(temp_dir, uploaded_file.name)

        with open(input_path, "wb") as f:
            f.write(uploaded_file.getvalue())

        subprocess.run(
            [
                r"C:\Program Files\LibreOffice\program\soffice.exe",
                "--headless",
                "--convert-to", "docx",
                "--outdir", temp_dir,
                input_path
            ],
            check=True,
            capture_output=True
        )

        output_path = os.path.join(
            temp_dir,
            os.path.splitext(uploaded_file.name)[0] + ".docx"
        )

        doc = Document(output_path)
        return "\n".join(p.text for p in doc.paragraphs)
st.set_page_config(
    page_title="ReadEase AI",
    page_icon="📖",
    layout="wide"
)

st.title("📖 ReadEase AI")
st.subheader("Your documents, read aloud.")

st.write(
    "Upload your PDF or Word document "
    "and listen to its original content."
)

uploaded_file = st.file_uploader(
    "Upload your document",
    type=["pdf", "docx", "doc"]
)

if uploaded_file:

    file_name = uploaded_file.name.lower()
    all_pages = []

    if file_name.endswith(".pdf"):

        doc = pymupdf.open(
            stream=uploaded_file.getvalue(),
            filetype="pdf"
        )

        for page in doc:
            all_pages.append(page.get_text())

        doc.close()

        st.success(
            f"PDF uploaded! Total pages: {len(all_pages)}"
        )

        page_number = st.selectbox(
            "Select a page",
            range(1, len(all_pages) + 1)
        )

        page_text = all_pages[page_number - 1]

    elif file_name.endswith(".doc"):
            page_text = read_doc_file(uploaded_file)
            st.success("Word document uploaded successfully!")

    else:
            doc = Document(uploaded_file)
            paragraphs = [
                p.text for p in doc.paragraphs if p.text.strip()
            ]
            tables = []
            for table in doc.tables:
                for row in table.rows:
                    tables.append(
                        " | ".join(cell.text for cell in row.cells)
                    )
            page_text = "\n".join(paragraphs + tables)
            st.success("Word document uploaded successfully!")

   
    if page_text.strip():

        st.subheader("Extracted Content")
        st.subheader("Listen to your document")

        safe_text = json.dumps(page_text).replace("<", "\\u003c")

        html = """
        <style>
        button {
            padding: 10px 16px;
            margin: 4px;
            border-radius: 8px;
            cursor: pointer;
        }
        #readingText {
            height: 300px;
            overflow-y: auto;
            padding: 15px;
            border: 1px solid #888;
            border-radius: 10px;
            line-height: 2;
            font-size: 17px;
        }
        .active {
            background-color: #ffe082;
            border-radius: 4px;
        }
        </style>

        <h3>Follow Along While Listening</h3>
        <div id="readingText"></div>

        <br>
        <button onclick="playText()">▶ Play</button>
        <button onclick="speechSynthesis.pause()">⏸ Pause</button>
        <button onclick="speechSynthesis.resume()">▶ Resume</button>
        <button onclick="stopText()">⏹ Stop</button>

        <br><br>
        <label>Select voice:</label>
        <select id="voiceSelect"></select>

        <br><br>
        <label>Reading speed:</label>
        <input id="speed" type="range"
        min="0.5" max="2" step="0.1" value="1">

        <script>
        const text = """ + safe_text + """;
        const container = document.getElementById("readingText");
        const voiceSelect = document.getElementById("voiceSelect");

        const sentences = text.match(/[^.!?]+[.!?]+\\s*|[^.!?]+$/g) || [text];

        let position = 0;
        const positions = [];

        sentences.forEach((sentence, index) => {
            positions.push(position);

            const span = document.createElement("span");
            span.className = "sentence";
            span.id = "sentence-" + index;
            span.textContent = sentence;
            container.appendChild(span);

            position += sentence.length;
        });

        function loadVoices() {
            const voices = speechSynthesis.getVoices();
            voiceSelect.innerHTML = "";

            voices.forEach((voice, index) => {
                const option = document.createElement("option");
                option.value = index;
                option.textContent = voice.name + " (" + voice.lang + ")";
                voiceSelect.appendChild(option);
            });
        }

        loadVoices();
        speechSynthesis.onvoiceschanged = loadVoices;

        function stopText() {
            speechSynthesis.cancel();
            document.querySelectorAll(".sentence").forEach(s =>
                s.classList.remove("active")
            );
        }

        function playText() {
            speechSynthesis.cancel();

            const speech = new SpeechSynthesisUtterance(text);
            const voices = speechSynthesis.getVoices();

            speech.rate = Number(document.getElementById("speed").value);

            if (voices[Number(voiceSelect.value)]) {
                speech.voice = voices[Number(voiceSelect.value)];
            }

            speech.onboundary = function(event) {
                if (event.name !== "word") return;

                let index = 0;

                for (let i = 0; i < positions.length; i++) {
                    if (event.charIndex >= positions[i]) {
                        index = i;
                    }
                }

                document.querySelectorAll(".sentence").forEach(s =>
                    s.classList.remove("active")
                );

                const current = document.getElementById("sentence-" + index);

                if (current) {
                    current.classList.add("active");
                    current.scrollIntoView({
                        behavior: "smooth",
                        block: "center"
                    });
                }
            };

            speechSynthesis.speak(speech);
        }
        </script>
        """

        components.html(html, height=530)

    else:
        st.warning(
            "No readable text found in this document."
        )
